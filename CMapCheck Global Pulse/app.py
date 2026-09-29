import os
import io
import re
import time as _time
from datetime import datetime, timezone
from flask import Flask, request, jsonify, render_template, send_file
from flask_cors import CORS
from flask_sqlalchemy import SQLAlchemy
from dotenv import load_dotenv
import google.generativeai as genai

# --- Report libraries ---
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.units import cm
from reportlab.lib import colors
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, HRFlowable
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor

load_dotenv()

app = Flask(__name__)
CORS(app)

# --- DATABASE ---
app.config['SQLALCHEMY_DATABASE_URI'] = 'postgresql://admin:password123@localhost:5432/world_data_db'
app.config['SQLALCHEMY_TRACK_MODIFICATIONS'] = False
db = SQLAlchemy(app)

class Country(db.Model):
    __tablename__ = 'countries'
    __table_args__ = {'extend_existing': True}
    country_name = db.Column(db.String(100), primary_key=True)
    country_code = db.Column(db.String(10))
    continent    = db.Column(db.String(50))
    flag_emoji   = db.Column(db.String(10))

# GeoJSON name → Vietnamese DB name
COUNTRY_NAME_MAP = {
    "Vietnam": "Việt Nam",
    "Thailand": "Thái Lan",
    "Indonesia": "Indonesia",
    "Malaysia": "Malaysia",
    "Philippines": "Philippines",
    "Singapore": "Singapore",
    "Laos": "Lào",
    "Cambodia": "Campuchia",
    "Myanmar": "Myanmar",
    "Brunei": "Brunei",
    "Russia": "Nga",
    "Ukraine": "Ukraine",
    "China": "Trung Quốc",
}

# --- GEMINI AI ---
def configure_gemini():
    from dotenv import load_dotenv
    load_dotenv(override=True)
    api_key = os.getenv("GEMINI_API_KEY")
    if api_key:
        genai.configure(api_key=api_key)
        global model
        model = genai.GenerativeModel('gemini-2.5-flash')

def generate_gemini_content(prompt):
    configure_gemini()
    try:
        return model.generate_content(prompt)
    except Exception as e:
        err_msg = str(e).lower()
        if any(x in err_msg for x in ["429", "quota", "limit", "exhausted"]):
            try:
                fallback = genai.GenerativeModel('gemini-flash-latest')
                return fallback.generate_content(prompt)
            except Exception as e2:
                fallback_lite = genai.GenerativeModel('gemini-2.5-flash-lite')
                return fallback_lite.generate_content(prompt)
        raise e

configure_gemini()

# --- AUTO-SCRAPE on startup ---
def _auto_scrape():
    import threading
    from scraper import run_all
    def _run():
        run_all(db_url=app.config['SQLALCHEMY_DATABASE_URI'], verbose=True)
    threading.Thread(target=_run, daemon=True).start()

with app.app_context():
    # Dọn dữ liệu rác trong DB khi khởi động
    try:
        from sqlalchemy import text as _t
        with db.engine.connect() as _conn:
            # Xóa bài báo cũ hơn 7 ngày
            _conn.execute(_t("DELETE FROM articles WHERE scraped_at < NOW() - INTERVAL '7 days'"))
            # Xóa duplicate articles (giữ bản mới nhất)
            _conn.execute(_t("""
                DELETE FROM articles a
                USING articles b
                WHERE a.id < b.id AND a.url = b.url
            """))
            # Xóa duplicate countries (giữ bản mới nhất)
            _conn.execute(_t("""
                DELETE FROM countries a
                USING countries b
                WHERE a.ctid < b.ctid AND a.country_name = b.country_name
            """))
            _conn.commit()
        print("[DB] Đã dọn dữ liệu rác")
    except Exception as _e:
        print(f"[DB] Cleanup skip: {_e}")
    _auto_scrape()

# ── ROUTES ────────────────────────────────────────────────────────────────────

@app.route('/')
def index():
    countries_from_db = Country.query.all()
    return render_template('index.html', countries=countries_from_db)

@app.route('/api/risk-data')
def get_risk_data():
    countries = Country.query.all()
    reverse_map = {v: k for k, v in COUNTRY_NAME_MAP.items()}
    result = []
    for c in countries:
        if not c.country_code:
            continue
        en_name = reverse_map.get(c.country_name, c.country_name)
        result.append({
            "code":       c.country_code,
            "name":       en_name,
            "name_vi":    c.country_name,
            "flag_emoji": c.flag_emoji,
        })
    return jsonify(result)

_summary_cache = {}

@app.route("/api/city-country-summary", methods=["POST"])
def get_city_country_summary():
    data    = request.json
    city    = data.get("city", "")
    country = data.get("country", "")

    # Thêm ngày vào cache_key để làm mới tóm tắt mỗi ngày
    from datetime import date
    cache_key = f"{country}|{city}|{date.today().isoformat()}"
    if cache_key in _summary_cache:
        return jsonify(_summary_cache[cache_key])

    # Lấy tin tức mới nhất từ DB làm bối cảnh để AI tóm tắt tình hình thật nhất
    news_context = ""
    headlines = []
    rows = []
    try:
        from sqlalchemy import text as sql_text
        name_map = {
            "Vietnam": "Vietnam", "Thailand": "Thailand", "Indonesia": "Indonesia",
            "Malaysia": "Malaysia", "Philippines": "Philippines", "Singapore": "Singapore",
            "Laos": "Laos", "Cambodia": "Cambodia", "Myanmar": "Myanmar", "Brunei": "Brunei",
            "Russia": "Russia", "Ukraine": "Ukraine", "China": "China",
        }
        country_en = name_map.get(country, country)
        
        with db.engine.connect() as conn:
            # Lấy toàn bộ bài viết trong vòng 7 ngày gần nhất để lọc thành phố chính xác đầy đủ
            rows = conn.execute(sql_text("""
                SELECT title, summary FROM articles 
                WHERE country_en = :country 
                  AND published_at >= NOW() - INTERVAL '7 days'
                ORDER BY published_at DESC NULLS LAST 
            """), {"country": country_en}).fetchall()
            
            if rows:
                if city:
                    CITY_KEYWORDS = {
                        "Hà Nội":    ["Hà Nội", "Hanoi"],
                        "TP.HCM":    ["TP.HCM", "Hồ Chí Minh", "Sài Gòn", "TPHCM"],
                        "Đà Nẵng":   ["Đà Nẵng", "Da Nang"],
                        "Hải Phòng": ["Hải Phòng"],
                        "Cần Thơ":   ["Cần Thơ"],
                        "Huế":       ["Huế", "Thừa Thiên"],
                        "Nha Trang": ["Nha Trang", "Khánh Hòa"],
                        "Biên Hòa":  ["Biên Hòa", "Đồng Nai"],
                        "Vũng Tàu":  ["Vũng Tàu", "Bà Rịa"],
                        "Quảng Ninh":["Quảng Ninh", "Hạ Long"],
                        "Hoàng Sa":  ["quần đảo hoàng sa", "huyện đảo hoàng sa", "paracel"],
                        "Trường Sa": ["quần đảo trường sa", "huyện đảo trường sa", "đảo trường sa", "spratly"],
                        "Bangkok":   ["Bangkok", "Băng Cốc"],
                        "Chiang Mai": ["Chiang Mai"],
                        "Phuket":    ["Phuket"],
                        "Jakarta":   ["Jakarta"],
                        "Surabaya":  ["Surabaya"],
                        "Bali":      ["Bali"],
                        "Kuala Lumpur": ["Kuala Lumpur", "KL"],
                        "Penang":    ["Penang"],
                        "Manila":    ["Manila"],
                        "Cebu":      ["Cebu"],
                        "Singapore": ["Singapore"],
                        "Phnom Penh": ["Phnom Penh"],
                        "Siem Reap": ["Siem Reap", "Angkor"],
                        "Yangon":    ["Yangon", "Rangoon"],
                        "Naypyidaw": ["Naypyidaw"],
                        "Vientiane": ["Vientiane"],
                        "Bandar Seri Begawan": ["Bandar", "Brunei", "Bandar Seri Begawan"],
                        "Moscow":    ["Moscow", "Moskva", "Mát-xcơ-va"],
                        "Saint Petersburg": ["Saint Petersburg", "St. Petersburg"],
                        "Kyiv":      ["Kyiv", "Kiev"],
                        "Kharkiv":   ["Kharkiv", "Kharkov"],
                        "Bắc Kinh":  ["Bắc Kinh", "Beijing"],
                        "Thượng Hải": ["Thượng Hải", "Shanghai"],
                    }
                    city_kws = CITY_KEYWORDS.get(city, []) if city else []
                    if city_kws:
                        valid_rows = [r for r in rows if r.title and any(kw.lower() in f"{r.title or ''} {r.summary or ''}".lower() for kw in city_kws)]
                    else:
                        valid_rows = [r for r in rows if r.title and city.lower() in f"{r.title or ''} {r.summary or ''}".lower()]
                else:
                    valid_rows = rows
                
                headlines = [r.title for r in valid_rows[:10] if r.title]
                news_lines = [f"- {r.title}" for r in valid_rows[:15] if r.title]
                news_context = "\n".join(news_lines)
    except Exception as e:
        print(f"[Error] Fetching summary news context failed: {e}")

    # Lấy thêm tin tức quốc gia để lồng ghép nếu cần
    national_headlines = [r.title for r in rows[:5] if r.title]

    if not city:
        if news_context:
            prompt = f"Dựa vào các tiêu đề tin tức trong 7 ngày gần nhất sau, hãy viết một đoạn văn ngắn (3-4 câu) kết nối các ý chặt chẽ, mô tả sinh động về tình hình chung HIỆN TẠI của quốc gia {country}. Tránh việc liệt kê khô khan, hãy dùng các câu liên kết tự nhiên. Ngôn ngữ: Tiếng Việt. Tuyệt đối không dùng gạch đầu dòng, không có câu dẫn (VD: 'Dưới đây là...', 'Dựa vào tin tức...'):\n{news_context}"
        else:
            prompt = f"Viết một đoạn văn ngắn (3 câu) mô tả sinh động tình hình hiện tại của quốc gia {country} trong 7 ngày qua. Ngôn ngữ: Tiếng Việt. Không gạch đầu dòng, không thêm câu dẫn."
    else:
        if news_context:
            prompt = f"Dựa vào các tin tức trong 7 ngày gần nhất sau, hãy viết một đoạn văn ngắn (3-4 câu) kết nối các ý chặt chẽ, mô tả sinh động về tình hình HIỆN TẠI của thành phố {city}, quốc gia {country}. Tránh việc liệt kê khô khan, hãy dùng các câu liên kết tự nhiên. Ngôn ngữ: Tiếng Việt. Tuyệt đối không dùng gạch đầu dòng, không có câu dẫn:\n{news_context}"
        else:
            # Nếu không có tin tức cho thành phố này, lồng ghép bối cảnh vĩ mô quốc gia
            national_context_str = "\n".join([f"- {h}" for h in national_headlines])
            if national_context_str:
                prompt = (
                    f"Thành phố {city} của quốc gia {country} trong 7 ngày qua duy trì ổn định và chưa ghi nhận vụ việc bất ổn cục bộ nào. "
                    f"Hãy viết một đoạn văn ngắn (2-3 câu) nhận định tình hình: khẳng định địa phương {city} vẫn an toàn ổn định, "
                    f"nhưng điểm qua 1-2 sự kiện nổi bật của quốc gia {country} dựa trên các tiêu đề sau để người dùng có cái nhìn toàn cảnh:\n{national_context_str}\n"
                    f"Ngôn ngữ: Tiếng Việt. Tuyệt đối không dùng gạch đầu dòng, không có câu dẫn."
                )
            else:
                prompt = f"Viết một đoạn văn ngắn (3 câu) mô tả sinh động tình hình của thành phố {city}, quốc gia {country} trong 7 ngày qua. Ngôn ngữ: Tiếng Việt. Không gạch đầu dòng, không thêm câu dẫn."
    
    try:
        response = generate_gemini_content(prompt)
        cleaned_text = re.sub(r'^(Dưới đây là|Theo các tin tức|Dựa vào|Tóm tắt|Sau đây là).*?:\s*', '', response.text.strip(), flags=re.IGNORECASE)
        result = {"summary": cleaned_text, "source": "gemini_dynamic"}
        _summary_cache[cache_key] = result
        return jsonify(result)
    except Exception as e:
        print(f"[Warning] Gemini API failed to generate summary: {e}. Using local heuristic premium fallback.")
        
        # Hàm sinh đoạn văn tóm tắt tối tân (Premium Heuristic Paragraph Generator)
        def generate_premium_heuristic_summary(headlines, place, country_name=None, nat_headlines=None):
            # Làm sạch tiêu đề để nối câu tự nhiên
            def clean_headline(h):
                h = h.strip()
                # Loại bỏ các dấu câu ở cuối
                h = re.sub(r'[.;:\-–|]+$', '', h).strip()
                # Chuyển chữ cái đầu thành viết thường nếu không phải danh từ riêng viết hoa
                words = h.split()
                if len(words) > 1:
                    first_word = words[0]
                    second_word = words[1]
                    # Nếu từ thứ hai không viết hoa chữ đầu, ta có thể viết thường từ đầu tiên
                    if second_word and second_word[0].islower() and not first_word.isupper():
                        h = h[0].lower() + h[1:]
                elif len(words) == 1:
                    h = h.lower()
                return h

            cleaned = [clean_headline(h) for h in headlines if h]
            
            # Trường hợp KHÔNG có tin riêng cho thành phố/địa phương này
            if not cleaned:
                if nat_headlines and country_name:
                    cleaned_nat = [clean_headline(h) for h in nat_headlines if h]
                    if len(cleaned_nat) >= 2:
                        return (
                            f"Trong 7 ngày qua, tình hình an ninh trật tự tại {place} duy trì ổn định và chưa ghi nhận vụ việc bất ổn cục bộ nào. "
                            f"Tuy nhiên, trên phạm vi cả nước ({country_name}), dư luận đang chú ý đến một số diễn biến nổi bật như {cleaned_nat[0]}, "
                            f"cùng với thông tin {cleaned_nat[1]}."
                        )
                    elif len(cleaned_nat) == 1:
                        return (
                            f"Trong 7 ngày qua, tình hình an ninh trật tự tại {place} duy trì ổn định, an toàn. "
                            f"Trên bình diện quốc gia ({country_name}), sự kiện đáng chú ý nhất gần đây là việc {cleaned_nat[0]}."
                        )
                return f"Trong 7 ngày qua, tình hình an ninh trật tự và đời sống dân sinh tại {place} duy trì ổn định, an toàn. Chưa ghi nhận vụ việc bất ổn, thiên tai hay biến động nghiêm trọng nào xảy ra trên địa bàn."

            if len(cleaned) == 1:
                return f"Tình hình hiện tại ở {place} ghi nhận diễn biến nổi bật liên quan đến việc {cleaned[0]}. Ngoài sự việc này, các hoạt động dân sinh, an ninh trật tự trên địa bàn vẫn duy trì ổn định bình thường."
                
            if len(cleaned) == 2:
                return f"Trong những ngày qua, {place} ghi nhận diễn biến đáng chú ý liên quan đến việc {cleaned[0]}. Bên cạnh đó, địa phương cũng có thông tin về việc {cleaned[1]}. Nhìn chung tình hình an ninh xã hội vẫn được kiểm soát tốt."
                
            # Từ 3 tin trở lên
            return (
                f"Trong 7 ngày qua, tình hình {place} ghi nhận những biến động và sự kiện đáng chú ý. "
                f"Nổi bật nhất là việc {cleaned[0]}. "
                f"Bên cạnh đó, khu vực cũng ghi nhận diễn biến quan trọng khác liên quan đến việc {cleaned[1]}. "
                f"Đồng thời, dư luận và truyền thông đang đặc biệt quan tâm theo dõi thông tin {cleaned[2]}."
            )

        summary_text = generate_premium_heuristic_summary(headlines[:3], city or country, country, national_headlines[:3])
        result = {"summary": summary_text, "source": "local_fallback"}
        _summary_cache[cache_key] = result
        return jsonify(result)


@app.route('/api/chat', methods=['POST'])
def chat_with_ai():
    data     = request.json
    message  = data.get('message', '')
    country  = data.get('country', '')
    realtime = data.get('realtime', '')
    
    # RAG: Lấy tin tức 7 ngày qua của quốc gia đó từ DB làm bối cảnh để AI trả lời chính xác, cập nhật nhất
    news_context = ""
    try:
        from sqlalchemy import text as sql_text
        name_map = {
            "Vietnam": "Vietnam", "Thailand": "Thailand", "Indonesia": "Indonesia",
            "Malaysia": "Malaysia", "Philippines": "Philippines", "Singapore": "Singapore",
            "Laos": "Laos", "Cambodia": "Cambodia", "Myanmar": "Myanmar", "Brunei": "Brunei",
            "Russia": "Russia", "Ukraine": "Ukraine", "China": "China",
        }
        country_en = name_map.get(country, country)
        with db.engine.connect() as conn:
            rows = conn.execute(sql_text("""
                SELECT title, summary, published_at FROM articles
                WHERE country_en = :country
                  AND published_at >= NOW() - INTERVAL '7 days'
                ORDER BY published_at DESC NULLS LAST
                LIMIT 30
            """), {"country": country_en}).fetchall()
            
            if rows:
                news_lines = []
                for r in rows:
                    pub_str = r.published_at.strftime('%d/%m/%Y %H:%M') if r.published_at else "Mới đây"
                    news_lines.append(f"- [{pub_str}] {r.title}: {r.summary or ''}")
                news_context = "\n".join(news_lines)
    except Exception as e:
        print(f"[Error] Failed to fetch news context for chat: {e}")

    prompt = (
        f"Bạn là trợ lý AI phân tích địa chính trị và dữ liệu toàn cầu của hệ thống CMapCheck.\n"
        f"Thời gian thực tế hiện tại: {realtime}.\n"
        f"Quốc gia đang xem xét: {country}.\n\n"
    )
    if news_context:
        prompt += (
            f"Dưới đây là danh sách các tin tức an ninh, biến động xã hội và sự kiện nổi bật MỚI NHẤT được hệ thống cào về trong 7 ngày gần đây tại {country} (Hãy ưu tiên sử dụng dữ liệu này để trả lời chính xác, cập nhật nhất về các sự việc thực tế):\n"
            f"{news_context}\n\n"
        )
    prompt += (
        f"Yêu cầu trả lời: Dựa vào thông tin bối cảnh tin tức ở trên (nếu có liên quan) kết hợp với kiến thức chuyên môn của bạn, trả lời ngắn gọn, khách quan, phân tích sâu sắc bằng tiếng Việt cho câu hỏi của người dùng.\n"
        f"Người dùng hỏi: {message}"
    )

    try:
        response = generate_gemini_content(prompt)
        return jsonify({"reply": response.text})
    except Exception as e:
        return jsonify({"error": str(e)}), 500

@app.route('/api/enhance-prompt', methods=['POST'])
def enhance_prompt():
    data    = request.json
    prompt  = data.get('prompt', '')
    style   = data.get('style', '')
    angle   = data.get('angle', '')
    country = data.get('country', '')
    instruction = (
        f"Translate and enhance this image prompt to detailed English for AI image generation. "
        f"Original: '{prompt}'. Country: {country}. Style: {style}. Angle: {angle}. "
        f"Return ONLY the enhanced English prompt, max 200 characters."
    )
    try:
        response = generate_gemini_content(instruction)
        return jsonify({"prompt": response.text.strip().strip('"').strip("'")})
    except Exception as e:
        return jsonify({"error": str(e)}), 500

@app.route('/api/generate-image', methods=['POST'])
def generate_image():
    import urllib.request, urllib.parse
    data       = request.json
    prompt     = data.get('prompt', '')
    resolution = data.get('resolution', '1024x1024')
    seed       = data.get('seed', 42)
    res_map    = {'1024x1024': (1024,1024), '1920x1080': (1920,1080), '4K UHD': (2048,2048)}
    w, h       = res_map.get(resolution, (1024, 1024))
    url        = f"https://image.pollinations.ai/prompt/{urllib.parse.quote(prompt)}?width={w}&height={h}&seed={seed}&nologo=true"
    try:
        req = urllib.request.Request(url, headers={'User-Agent': 'Mozilla/5.0'})
        with urllib.request.urlopen(req, timeout=15) as r:
            return send_file(io.BytesIO(r.read()), mimetype=r.headers.get('Content-Type', 'image/jpeg'))
    except Exception as e:
        return jsonify({"error": str(e)}), 500

@app.route('/api/news', methods=['POST'])
def get_news():
    """Lấy tin tức mới nhất theo quốc gia từ DB."""
    from sqlalchemy import text as sql_text
    data        = request.json or {}
    country_en  = data.get('country', '')
    limit       = int(data.get('limit', 10))

    name_map = {
        "Vietnam": "Vietnam", "Thailand": "Thailand", "Indonesia": "Indonesia",
        "Malaysia": "Malaysia", "Philippines": "Philippines", "Singapore": "Singapore",
        "Laos": "Laos", "Cambodia": "Cambodia", "Myanmar": "Myanmar", "Brunei": "Brunei",
        "Russia": "Russia", "Ukraine": "Ukraine", "China": "China",
    }
    country_en = name_map.get(country_en, country_en)

    try:
        with db.engine.connect() as conn:
            rows = conn.execute(sql_text("""
                SELECT title, summary, url, source, country_name, published_at
                FROM articles
                WHERE country_en = :country
                ORDER BY published_at DESC NULLS LAST
                LIMIT :limit
            """), {"country": country_en, "limit": limit}).fetchall()

        articles = []
        for r in rows:
            pub = r.published_at
            if pub:
                pub_utc  = pub.replace(tzinfo=timezone.utc) if pub.tzinfo is None else pub
                diff_h   = int((datetime.now(timezone.utc) - pub_utc).total_seconds() / 3600)
                time_str = "Vừa xong" if diff_h < 1 else (f"{diff_h} giờ trước" if diff_h < 24 else f"{diff_h//24} ngày trước")
            else:
                time_str = ""
            articles.append({"title": r.title, "summary": r.summary,
                              "url": r.url, "source": r.source, "time": time_str})
        return jsonify({"articles": articles, "country": country_en})
    except Exception as e:
        return jsonify({"articles": [], "error": str(e)})


@app.route('/api/city-data', methods=['POST'])
def get_city_data():
    """Trả về danh sách thành phố với số bài báo và prominence score."""
    from sqlalchemy import text as sql_text
    data       = request.json or {}
    country    = data.get('country', '')
    # Mapping tên GeoJSON → tên tiếng Việt trong DB
    name_map = {
        "Vietnam": "Việt Nam", "Thailand": "Thái Lan", "Indonesia": "Indonesia",
        "Malaysia": "Malaysia", "Philippines": "Philippines", "Singapore": "Singapore",
        "Laos": "Lào", "Cambodia": "Campuchia", "Myanmar": "Myanmar", "Brunei": "Brunei",
        "Russia": "Nga", "Ukraine": "Ukraine", "China": "Trung Quốc",
    }
    country_vi = name_map.get(country, country)

    # Tọa độ các thành phố chính theo quốc gia
    CITY_COORDS = {
        "Việt Nam": [
            {"name": "Hà Nội",       "lat": 21.0285, "lng": 105.8542, "keywords": ["Hà Nội", "Hanoi"]},
            {"name": "TP.HCM",       "lat": 10.8231, "lng": 106.6297, "keywords": ["TP.HCM", "Hồ Chí Minh", "Sài Gòn", "TPHCM"]},
            {"name": "Đà Nẵng",      "lat": 16.0544, "lng": 108.2022, "keywords": ["Đà Nẵng", "Da Nang"]},
            {"name": "Hải Phòng",    "lat": 20.8449, "lng": 106.6881, "keywords": ["Hải Phòng"]},
            {"name": "Cần Thơ",      "lat": 10.0452, "lng": 105.7469, "keywords": ["Cần Thơ"]},
            {"name": "Huế",          "lat": 16.4637, "lng": 107.5909, "keywords": ["Huế", "Thừa Thiên"]},
            {"name": "Nha Trang",    "lat": 12.2388, "lng": 109.1967, "keywords": ["Nha Trang", "Khánh Hòa"]},
            {"name": "Biên Hòa",     "lat": 10.9574, "lng": 106.8426, "keywords": ["Biên Hòa", "Đồng Nai"]},
            {"name": "Vũng Tàu",     "lat": 10.3460, "lng": 107.0843, "keywords": ["Vũng Tàu", "Bà Rịa"]},
            {"name": "Quảng Ninh",   "lat": 21.0064, "lng": 107.2925, "keywords": ["Quảng Ninh", "Hạ Long"]},
            {"name": "Hoàng Sa",     "lat": 16.3986, "lng": 112.0163, "keywords": ["Hoàng Sa", "Paracel"]},
            {"name": "Trường Sa",    "lat": 10.0245, "lng": 114.1848, "keywords": ["Trường Sa", "Spratly"]},
        ],
        "Thái Lan": [
            {"name": "Bangkok",      "lat": 13.7563, "lng": 100.5018, "keywords": ["Bangkok", "Băng Cốc"]},
            {"name": "Chiang Mai",   "lat": 18.7883, "lng": 98.9853,  "keywords": ["Chiang Mai"]},
            {"name": "Phuket",       "lat": 7.8804,  "lng": 98.3923,  "keywords": ["Phuket"]},
        ],
        "Indonesia": [
            {"name": "Jakarta",      "lat": -6.2088, "lng": 106.8456, "keywords": ["Jakarta"]},
            {"name": "Surabaya",     "lat": -7.2575, "lng": 112.7521, "keywords": ["Surabaya"]},
            {"name": "Bali",         "lat": -8.3405, "lng": 115.0920, "keywords": ["Bali"]},
        ],
        "Malaysia": [
            {"name": "Kuala Lumpur", "lat": 3.1390,  "lng": 101.6869, "keywords": ["Kuala Lumpur", "KL"]},
            {"name": "Penang",       "lat": 5.4141,  "lng": 100.3288, "keywords": ["Penang"]},
        ],
        "Philippines": [
            {"name": "Manila",       "lat": 14.5995, "lng": 120.9842, "keywords": ["Manila"]},
            {"name": "Cebu",         "lat": 10.3157, "lng": 123.8854, "keywords": ["Cebu"]},
        ],
        "Singapore": [
            {"name": "Singapore",    "lat": 1.3521,  "lng": 103.8198, "keywords": ["Singapore"]},
        ],
        "Campuchia": [
            {"name": "Phnom Penh",   "lat": 11.5564, "lng": 104.9282, "keywords": ["Phnom Penh"]},
            {"name": "Siem Reap",    "lat": 13.3671, "lng": 103.8448, "keywords": ["Siem Reap", "Angkor"]},
        ],
        "Myanmar": [
            {"name": "Yangon",       "lat": 16.8661, "lng": 96.1951,  "keywords": ["Yangon", "Rangoon"]},
            {"name": "Naypyidaw",    "lat": 19.7633, "lng": 96.0785,  "keywords": ["Naypyidaw"]},
        ],
        "Lào": [
            {"name": "Vientiane",    "lat": 17.9757, "lng": 102.6331, "keywords": ["Vientiane"]},
        ],
        "Brunei": [
            {"name": "Bandar Seri Begawan", "lat": 4.9031, "lng": 114.9398, "keywords": ["Bandar", "Brunei"]},
        ],
        "Nga": [
            {"name": "Moscow", "lat": 55.7558, "lng": 37.6173, "keywords": ["Moscow", "Moskva", "Mát-xcơ-va"]},
            {"name": "Saint Petersburg", "lat": 59.9311, "lng": 30.3609, "keywords": ["Saint Petersburg", "St. Petersburg"]},
        ],
        "Ukraine": [
            {"name": "Kyiv", "lat": 50.4501, "lng": 30.5234, "keywords": ["Kyiv", "Kiev"]},
            {"name": "Kharkiv", "lat": 49.9935, "lng": 36.2304, "keywords": ["Kharkiv", "Kharkov"]},
        ],
        "Trung Quốc": [
            {"name": "Bắc Kinh", "lat": 39.9042, "lng": 116.4074, "keywords": ["Bắc Kinh", "Beijing"]},
            {"name": "Thượng Hải", "lat": 31.2304, "lng": 121.4737, "keywords": ["Thượng Hải", "Shanghai"]},
        ],
    }

    cities_config = CITY_COORDS.get(country_vi, [])
    if not cities_config:
        return jsonify({"cities": []})

    try:
        # Lấy tất cả bài báo của quốc gia trong 7 ngày
        with db.engine.connect() as conn:
            from sqlalchemy import text as sql_text
            rows = conn.execute(sql_text("""
                SELECT title, summary
                FROM articles
                WHERE country_name = :country
                  AND scraped_at >= NOW() - INTERVAL '7 days'
            """), {"country": country_vi}).fetchall()

        # Gộp toàn bộ text để tìm kiếm
        all_text = " ".join(
            f"{r.title or ''} {r.summary or ''}" for r in rows
        ).lower()
        total_articles = max(len(rows), 1)

        RISK_WORDS = ["bão", "lũ", "ngập", "tai nạn", "cháy", "nổ", "sạt lở", "động đất", "thiên tai", "biểu tình", "bạo loạn", "xung đột", "chiến tranh", "khủng bố", "bắt giữ", "khởi tố", "án mạng", "lừa đảo", "tội phạm", "ma túy", "tham nhũng", "kỷ luật", "cảnh báo", "nguy hiểm", "căng thẳng", "thiệt hại", "tử vong", "thương vong", "bệnh dịch", "suy thoái", "cảnh sát", "điều tra", "thiệt mạng"]
        STABLE_WORDS = ["phát triển", "khai mạc", "đầu tư", "kỷ niệm", "lễ hội", "hòa bình", "tăng trưởng", "hợp tác", "thành công", "cải thiện", "xây dựng", "khánh thành", "văn hóa", "du lịch", "tích cực", "hỗ trợ", "tôn vinh"]

        cities_result = []
        for city in cities_config:
            count = 0
            risk_points = 0
            for r in rows:
                text = f"{r.title or ''} {r.summary or ''}".lower()
                if any(kw.lower() in text for kw in city["keywords"]):
                    count += 1
                    has_risk = any(rw in text for rw in RISK_WORDS)
                    has_stable = any(sw in text for sw in STABLE_WORDS)
                    
                    if has_risk and not has_stable:
                        risk_points += 35
                    elif has_risk and has_stable:
                        risk_points += 15
                    elif has_stable:
                        risk_points -= 10
                    else:
                        risk_points += 5
            
            if count == 0:
                prominence = 0 # Không có dữ liệu -> xám
            else:
                score = 15 + risk_points
                prominence = max(5, min(int(score), 100))

            cities_result.append({
                "name":       city["name"],
                "lat":        city["lat"],
                "lng":        city["lng"],
                "prominence": prominence,
                "news_count": count,
            })

        # Sắp xếp theo prominence giảm dần
        cities_result.sort(key=lambda x: -x["prominence"])
        return jsonify({"cities": cities_result, "country": country})

    except Exception as e:
        # Fallback: trả về thành phố với prominence = 0 nếu DB lỗi
        fallback = [
            {"name": c["name"], "lat": c["lat"], "lng": c["lng"],
             "prominence": 0, "news_count": 0}
            for c in cities_config
        ]
        return jsonify({"cities": fallback, "error": str(e)})


@app.route('/api/news-data', methods=['POST'])
def get_news_data():
    """Lấy tin tức theo quốc gia + ngày + thành phố (nếu có). Có cache 60s."""
    from sqlalchemy import text as sql_text
    from datetime import timedelta as _td
    import time as _t
    data    = request.json or {}
    country = data.get('country', '')
    city    = data.get('city', '').strip()
    date    = data.get('date', '')

    # ── Cache key ──
    cache_key = f"{country}|{date}|{city}"
    now_ts = _t.time()
    if cache_key in _news_cache:
        cached_data, cached_ts = _news_cache[cache_key]
        if now_ts - cached_ts < 60:  # cache 60 giây
            return jsonify(cached_data)

    name_map = {
        "Vietnam": "Vietnam", "Thailand": "Thailand", "Indonesia": "Indonesia",
        "Malaysia": "Malaysia", "Philippines": "Philippines", "Singapore": "Singapore",
        "Laos": "Laos", "Cambodia": "Cambodia", "Myanmar": "Myanmar", "Brunei": "Brunei",
        "Russia": "Russia", "Ukraine": "Ukraine", "China": "China",
    }
    country_en = name_map.get(country, country)

    # Keywords của từng thành phố để lọc phía Python (nhanh hơn LIKE nhiều cột)
    CITY_KEYWORDS = {
        "Hà Nội":    ["Hà Nội", "Hanoi", "hà nội"],
        "TP.HCM":    ["TP.HCM", "Hồ Chí Minh", "Sài Gòn", "TPHCM", "tp.hcm", "hồ chí minh"],
        "Đà Nẵng":   ["Đà Nẵng", "Da Nang", "đà nẵng"],
        "Hải Phòng": ["Hải Phòng", "hải phòng"],
        "Cần Thơ":   ["Cần Thơ", "cần thơ"],
        "Huế":       ["Huế", "Thừa Thiên", "huế"],
        "Nha Trang": ["Nha Trang", "Khánh Hòa", "nha trang"],
        "Biên Hòa":  ["Biên Hòa", "Đồng Nai", "biên hòa"],
        "Vũng Tàu":  ["Vũng Tàu", "Bà Rịa", "vũng tàu"],
        "Quảng Ninh":["Quảng Ninh", "Hạ Long", "quảng ninh"],
        "Hoàng Sa":  ["quần đảo hoàng sa", "huyện đảo hoàng sa", "paracel"],
        "Trường Sa": ["quần đảo trường sa", "huyện đảo trường sa", "đảo trường sa", "spratly"],
        "Bangkok":   ["Bangkok", "Băng Cốc"],
        "Jakarta":   ["Jakarta"],
        "Manila":    ["Manila"],
        "Singapore": ["Singapore"],
        "Kuala Lumpur": ["Kuala Lumpur", "KL"],
    }

    try:
        with db.engine.connect() as conn:
            # Xây dựng query theo date
            if date and len(date) == 10:
                try:
                    date_obj   = datetime.strptime(date, '%Y-%m-%d')
                    start_date = date_obj.replace(hour=0, minute=0, second=0, microsecond=0)
                    end_date   = start_date + _td(days=1)
                    rows = conn.execute(sql_text("""
                        SELECT title, summary, url, source, country_name, published_at
                        FROM articles
                        WHERE country_en = :country
                          AND published_at >= :start_date
                          AND published_at < :end_date
                        ORDER BY published_at DESC NULLS LAST
                        LIMIT 60
                    """), {"country": country_en,
                           "start_date": start_date,
                           "end_date":   end_date}).fetchall()
                except ValueError:
                    rows = []
            else:
                rows = conn.execute(sql_text("""
                    SELECT title, summary, url, source, country_name, published_at
                    FROM articles
                    WHERE country_en = :country
                    ORDER BY published_at DESC NULLS LAST
                    LIMIT 60
                """), {"country": country_en}).fetchall()

        # Lọc theo thành phố phía Python
        city_kws = CITY_KEYWORDS.get(city, []) if city else []

        # Nếu city được chỉ định nhưng không có trong map → không có tin
        if city and not city_kws:
            return jsonify({"news": [], "keywords": [], "country": country, "city": city})

        filtered = []
        for r in rows:
            if city_kws:
                text = f"{r.title or ''} {r.summary or ''}".lower()
                if not any(kw.lower() in text for kw in city_kws):
                    continue
            filtered.append(r)
            if len(filtered) >= 15:
                break

        news = []
        for r in filtered:
            pub = r.published_at
            if pub:
                pub_utc  = pub.replace(tzinfo=timezone.utc) if pub.tzinfo is None else pub
                diff_h   = int((datetime.now(timezone.utc) - pub_utc).total_seconds() / 3600)
                time_str = ("Vừa xong" if diff_h < 1
                            else f"{diff_h} giờ trước" if diff_h < 24
                            else f"{diff_h//24} ngày trước")
            else:
                time_str = ""
            news.append({
                "title":   r.title,
                "summary": r.summary,
                "url":     r.url,
                "source":  r.source,
                "city":    city,
                "time":    time_str,
                "prominence": 50,
            })

        result = {"news": news, "keywords": [], "country": country, "city": city}
        _news_cache[cache_key] = (result, now_ts)
        # Dọn cache cũ (giữ tối đa 200 entries)
        if len(_news_cache) > 200:
            oldest = sorted(_news_cache, key=lambda k: _news_cache[k][1])[:50]
            for k in oldest:
                del _news_cache[k]
        return jsonify(result)
    except Exception as e:
        return jsonify({"news": [], "keywords": [], "error": str(e)})


@app.route("/api/keywords", methods=["POST"])
def get_keywords():
    data       = request.json or {}
    country_en = data.get("country", "")
    city       = data.get("city", "").strip()
    limit      = int(data.get("limit", 60))

    name_map = {
        "Vietnam": "Vietnam", "Thailand": "Thailand", "Indonesia": "Indonesia",
        "Malaysia": "Malaysia", "Philippines": "Philippines", "Singapore": "Singapore",
        "Laos": "Laos", "Cambodia": "Cambodia", "Myanmar": "Myanmar", "Brunei": "Brunei",
        "Russia": "Russia", "Ukraine": "Ukraine", "China": "China",
    }
    country_en = name_map.get(country_en, country_en)

    # City keywords map (giống get_news_data)
    CITY_KEYWORDS = {
        "Hà Nội":    ["Hà Nội", "Hanoi"],
        "TP.HCM":    ["TP.HCM", "Hồ Chí Minh", "Sài Gòn", "TPHCM"],
        "Đà Nẵng":   ["Đà Nẵng", "Da Nang"],
        "Hải Phòng": ["Hải Phòng"],
        "Cần Thơ":   ["Cần Thơ"],
        "Huế":       ["Huế", "Thừa Thiên"],
        "Nha Trang": ["Nha Trang", "Khánh Hòa"],
        "Biên Hòa":  ["Biên Hòa", "Đồng Nai"],
        "Vũng Tàu":  ["Vũng Tàu", "Bà Rịa"],
        "Quảng Ninh":["Quảng Ninh", "Hạ Long"],
        "Hoàng Sa":  ["quần đảo hoàng sa", "huyện đảo hoàng sa", "paracel"],
        "Trường Sa": ["quần đảo trường sa", "huyện đảo trường sa", "đảo trường sa", "spratly"],
        "Bangkok":   ["Bangkok", "Băng Cốc"],
        "Chiang Mai": ["Chiang Mai"],
        "Phuket":    ["Phuket"],
        "Jakarta":   ["Jakarta"],
        "Surabaya":  ["Surabaya"],
        "Bali":      ["Bali"],
        "Kuala Lumpur": ["Kuala Lumpur", "KL"],
        "Penang":    ["Penang"],
        "Manila":    ["Manila"],
        "Cebu":      ["Cebu"],
        "Singapore": ["Singapore"],
        "Phnom Penh": ["Phnom Penh"],
        "Siem Reap": ["Siem Reap", "Angkor"],
        "Yangon":    ["Yangon", "Rangoon"],
        "Naypyidaw": ["Naypyidaw"],
        "Vientiane": ["Vientiane"],
        "Bandar Seri Begawan": ["Bandar", "Brunei", "Bandar Seri Begawan"],
        "Moscow":    ["Moscow", "Moskva", "Mát-xcơ-va"],
        "Saint Petersburg": ["Saint Petersburg", "St. Petersburg"],
        "Kyiv":      ["Kyiv", "Kiev"],
        "Kharkiv":   ["Kharkiv", "Kharkov"],
        "Bắc Kinh":  ["Bắc Kinh", "Beijing"],
        "Thượng Hải": ["Thượng Hải", "Shanghai"],
    }
    city_kws = CITY_KEYWORDS.get(city, []) if city else []

    # Định nghĩa 11 danh mục bất ổn và biến động xã hội lớn
    categories_definition = {
        "Dịch bệnh & Y tế": ["covid", "corona", "dịch bệnh", "y tế", "vắc xin", "sốt xuất huyết", "sởi", "lây nhiễm", "virus", "cúm", "bệnh nhân", "dịch sởi", "dịch cúm"],
        "Buôn lậu & Ma túy": ["ma túy", "heroin", "chất cấm", "buôn lậu", "hàng lậu", "vận chuyển trái phép", "pháo lậu", "vàng lậu", "thuốc lá lậu", "chất ma túy"],
        "Cướp giật & Trộm cắp": ["cướp", "trộm", "cướp giật", "trộm cắp", "bắt cóc", "tống tiền", "giật đồ", "đột nhập", "cướp tiệm vàng", "cướp tài sản"],
        "Bạo động & Mất an ninh": ["bạo động", "biểu tình", "ẩu đả", "hỗn chiến", "giang hồ", "gây rối", "vũ khí", "súng", "đâm chém", "an ninh trật tự", "gây rối trật tự"],
        "Thiên tai & Bão lũ": ["bão", "lũ", "ngập lụt", "sạt lở", "triều cường", "hạn hán", "động đất", "giông lốc", "thiên tai", "mưa lớn", "triều cường"],
        "Cháy nổ nghiêm trọng": ["cháy", "hỏa hoạn", "vụ nổ", "nổ lớn", "bình gas", "chập điện", "cứu hỏa", "đám cháy"],
        "Kỷ luật & Tham nhũng": ["kỷ luật", "khai trừ", "cách chức", "tham nhũng", "nhận hối lộ", "sai phạm", "bắt giam cán bộ", "tòa án", "xét xử", "khởi tố", "truy tố", "vụ án tham nhũng"],
        "Biến động giá cả & Vàng & USD": ["giá vàng", "tỷ giá", "usd", "lạm phát", "giá xăng", "tăng giá", "lãi suất", "ngân hàng miếng", "giá cả", "sjc"],
        "Tai nạn giao thông": ["tai nạn", "va chạm", "tông nhau", "lật xe", "xe khách", "tử vong đường bộ", "tai nạn giao thông", "đâm nhau"],
        "Tranh chấp & Lãnh hải": ["biển đông", "lãnh hải", "hoàng sa", "trường sa", "xâm phạm", "tàu nước ngoài", "quân sự", "đối ngoại", "đường chín đoạn"],
        "Trốn thuế & Lừa đảo tài chính": ["lừa đảo", "chiếm đoạt", "đa cấp", "trốn thuế", "giả danh", "hồ sơ giả", "giấy phép giả", "tài chính", "mạo danh", "lừa đảo qua mạng"]
    }

    try:
        from sqlalchemy import text as sql_text
        with db.engine.connect() as conn:
            # Luôn lấy bài báo trong vòng 7 ngày gần nhất, không bị giới hạn quá sớm để lọc được đầy đủ dữ liệu thành phố
            rows = conn.execute(sql_text("""
                SELECT title, summary, url, source
                FROM articles
                WHERE country_en = :country
                  AND published_at >= NOW() - INTERVAL '7 days'
                ORDER BY published_at DESC NULLS LAST
            """), {"country": country_en}).fetchall()

        # Lọc theo city nếu có
        if city_kws:
            filtered = [r for r in rows
                        if any(kw.lower() in f"{r.title or ''} {r.summary or ''}".lower()
                               for kw in city_kws)]
        else:
            filtered = list(rows)

        if not filtered:
            return jsonify({"keywords": [], "country": country_en, "city": city})

        keywords = []
        try:
            # Chuẩn bị nội dung gửi AI phân tích
            all_content = "\n".join([
                f"Tiêu đề: {r.title}" for r in filtered[:40] if r.title
            ])
            # Gọi Gemini phân loại các chủ đề biến động an ninh xã hội đang xảy ra thực tế
            prompt = (
                f"Dựa trên các tiêu đề tin tức sau của {city or country_en} trong 7 ngày qua, "
                f"hãy phân tích xem khu vực này đang xảy ra những biến động, sự cố, hoặc vấn đề bất ổn nào nổi bật nhất mà người dân quan tâm. "
                f"Hãy chọn và phân loại chúng vào các nhóm chính xác sau nếu thực sự có tin tức tương ứng:\n"
                f"- 'Dịch bệnh & Y tế'\n"
                f"- 'Buôn lậu & Ma túy'\n"
                f"- 'Cướp giật & Trộm cắp'\n"
                f"- 'Bạo động & Mất an ninh'\n"
                f"- 'Thiên tai & Bão lũ'\n"
                f"- 'Cháy nổ nghiêm trọng'\n"
                f"- 'Kỷ luật & Tham nhũng'\n"
                f"- 'Biến động giá cả & Vàng & USD'\n"
                f"- 'Tai nạn giao thông'\n"
                f"- 'Tranh chấp & Lãnh hải'\n"
                f"- 'Trốn thuế & Lừa đảo tài chính'\n\n"
                f"Chỉ trả về JSON array chứa tên các chủ đề đang thực sự xuất hiện trong tin tức (không bịa đặt thêm). "
                f"Chỉ trả về JSON array thuần túy tiếng Việt, không markdown, không giải thích:\n\n"
                f"{all_content}\n\n"
                f'Ví dụ output: ["Dịch bệnh & Y tế", "Buôn lậu & Ma túy", "Cháy nổ nghiêm trọng"]'
            )
            response = generate_gemini_content(prompt)
            raw = response.text.strip()

            import json as _json, re as _re
            m = _re.search(r'\[.*?\]', raw, _re.DOTALL)
            if m:
                extracted = _json.loads(m.group())
                # Chỉ lấy các danh mục hợp lệ nằm trong định nghĩa
                keywords = [k for k in extracted if k in categories_definition]
        except Exception as ai_err:
            print(f"[Warning] Gemini keyword classification failed: {ai_err}. Falling back to regex rule.")

        # Thuật toán cục bộ phân loại dựa trên Regex từ khóa (chạy khi Gemini lỗi hoặc không nhận diện được)
        if not keywords:
            category_counts = {}
            for cat, keywords_list in categories_definition.items():
                match_count = 0
                for r in filtered:
                    text_lower = f"{r.title or ''} {r.summary or ''}".lower()
                    if any(kw in text_lower for kw in keywords_list):
                        match_count += 1
                if match_count > 0:
                    category_counts[cat] = match_count
            
            # Sắp xếp các danh mục biến động có nhiều bài báo nhất lên trên
            sorted_cats = sorted(category_counts.items(), key=lambda x: x[1], reverse=True)
            keywords = [cat for cat, count in sorted_cats[:10]]

        # Map mỗi danh mục biến động -> danh sách các bài báo thực tế chứng minh
        keyword_articles = {}
        for cat in keywords:
            related = []
            keywords_list = categories_definition.get(cat, [])
            for r in filtered:
                text_lower = f"{r.title or ''} {r.summary or ''}".lower()
                if any(kw in text_lower for kw in keywords_list):
                    related.append({
                        "title":  r.title,
                        "url":    r.url,
                        "source": r.source,
                    })
            if related:
                keyword_articles[cat] = related[:5]

        return jsonify({
            "keywords":         keywords,
            "keyword_articles": keyword_articles,
            "country":          country_en,
            "city":             city,
        })
    except Exception as e:
        print(f"[Error] get_keywords overall error: {e}")
        return jsonify({
            "keywords":         [],
            "keyword_articles": {},
            "country":          country_en,
            "city":             city,
        })


# ── RISK SCORES ───────────────────────────────────────────────────────────────
_risk_cache      = {}
_risk_cache_time = 0
_news_cache      = {}  # key: "country|date|city" → (result, timestamp)

@app.route('/api/risk-scores')
def get_risk_scores():
    global _risk_cache, _risk_cache_time
    if _risk_cache and (_time.time() - _risk_cache_time) < 3600:
        return jsonify(_risk_cache)

    prompt = """You are a geopolitical analyst. Rate the current risk/instability level (2026) for each country.
Scale: 0=fully stable, 100=severe conflict/crisis.
Criteria: armed conflict, political instability, economic crisis, terrorism, human rights.
Return ONLY a valid JSON object, no markdown, no explanation:
{"Afghanistan":95,"Albania":15,"Algeria":35,"Angola":40,"Argentina":45,"Armenia":55,"Australia":5,"Austria":5,"Azerbaijan":50,"Bahrain":30,"Bangladesh":50,"Belarus":65,"Belgium":8,"Bolivia":40,"Bosnia and Herzegovina":35,"Brazil":45,"Brunei":10,"Bulgaria":20,"Burkina Faso":85,"Cambodia":30,"Cameroon":60,"Canada":5,"Central African Republic":90,"Chad":80,"Chile":25,"China":40,"Colombia":55,"Congo":65,"Costa Rica":10,"Croatia":15,"Cuba":60,"Czech Republic":8,"Democratic Republic of the Congo":88,"Denmark":3,"Dominican Republic":40,"Ecuador":50,"Egypt":55,"El Salvador":45,"Ethiopia":75,"Finland":3,"France":20,"Germany":8,"Ghana":20,"Greece":25,"Guatemala":50,"Guinea":55,"Haiti":85,"Honduras":55,"Hungary":25,"India":40,"Indonesia":30,"Iran":70,"Iraq":75,"Ireland":5,"Israel":80,"Italy":18,"Ivory Coast":40,"Jamaica":45,"Japan":5,"Jordan":35,"Kazakhstan":30,"Kenya":45,"Kosovo":35,"Kuwait":20,"Kyrgyzstan":40,"Laos":25,"Lebanon":75,"Libya":80,"Malaysia":15,"Mali":85,"Mexico":55,"Moldova":40,"Mongolia":20,"Morocco":30,"Mozambique":60,"Myanmar":85,"Nepal":35,"Netherlands":5,"Nicaragua":55,"Niger":80,"Nigeria":70,"North Korea":75,"Norway":3,"Oman":20,"Pakistan":75,"Palestine":95,"Panama":20,"Papua New Guinea":55,"Paraguay":30,"Peru":45,"Philippines":45,"Poland":20,"Portugal":8,"Qatar":15,"Romania":20,"Russia":80,"Rwanda":35,"Saudi Arabia":35,"Senegal":30,"Serbia":30,"Sierra Leone":45,"Singapore":3,"Somalia":90,"South Africa":50,"South Korea":15,"South Sudan":90,"Spain":15,"Sri Lanka":35,"Sudan":85,"Sweden":5,"Switzerland":3,"Syria":90,"Taiwan":45,"Tajikistan":45,"Tanzania":30,"Thailand":35,"Tunisia":35,"Turkey":50,"Turkmenistan":50,"Uganda":45,"Ukraine":90,"United Arab Emirates":15,"United Kingdom":15,"United States":25,"Uruguay":10,"Uzbekistan":35,"Venezuela":75,"Vietnam":20,"Yemen":90,"Zambia":35,"Zimbabwe":55}
Update the numbers based on actual 2026 situation and return the complete JSON."""

    try:
        response  = generate_gemini_content(prompt)
        text      = response.text.strip()
        json_match = re.search(r'\{[\s\S]+\}', text)
        if json_match:
            import json as _json
            scores = _json.loads(json_match.group())
            _risk_cache      = scores
            _risk_cache_time = _time.time()
            return jsonify(scores)
        raise ValueError("Could not parse JSON from Gemini response")
    except Exception as e:
        print(f"[Warning] Gemini API failed to get risk scores: {e}. Falling back to default data.")
        # Fallback to local default scores to prevent map from breaking
        fallback_scores = {"Afghanistan":95,"Albania":15,"Algeria":35,"Angola":40,"Argentina":45,"Armenia":55,"Australia":5,"Austria":5,"Azerbaijan":50,"Bahrain":30,"Bangladesh":50,"Belarus":65,"Belgium":8,"Bolivia":40,"Bosnia and Herzegovina":35,"Brazil":45,"Brunei":10,"Bulgaria":20,"Burkina Faso":85,"Cambodia":30,"Cameroon":60,"Canada":5,"Central African Republic":90,"Chad":80,"Chile":25,"China":40,"Colombia":55,"Congo":65,"Costa Rica":10,"Croatia":15,"Cuba":60,"Czech Republic":8,"Democratic Republic of the Congo":88,"Denmark":3,"Dominican Republic":40,"Ecuador":50,"Egypt":55,"El Salvador":45,"Ethiopia":75,"Finland":3,"France":20,"Germany":8,"Ghana":20,"Greece":25,"Guatemala":50,"Guinea":55,"Haiti":85,"Honduras":55,"Hungary":25,"India":40,"Indonesia":30,"Iran":70,"Iraq":75,"Ireland":5,"Israel":80,"Italy":18,"Ivory Coast":40,"Jamaica":45,"Japan":5,"Jordan":35,"Kazakhstan":30,"Kenya":45,"Kosovo":35,"Kuwait":20,"Kyrgyzstan":40,"Laos":25,"Lebanon":75,"Libya":80,"Malaysia":15,"Mali":85,"Mexico":55,"Moldova":40,"Mongolia":20,"Morocco":30,"Mozambique":60,"Myanmar":85,"Nepal":35,"Netherlands":5,"Nicaragua":55,"Niger":80,"Nigeria":70,"North Korea":75,"Norway":3,"Oman":20,"Pakistan":75,"Palestine":95,"Panama":20,"Papua New Guinea":55,"Paraguay":30,"Peru":45,"Philippines":45,"Poland":20,"Portugal":8,"Qatar":15,"Romania":20,"Russia":80,"Rwanda":35,"Saudi Arabia":35,"Senegal":30,"Serbia":30,"Sierra Leone":45,"Singapore":3,"Somalia":90,"South Africa":50,"South Korea":15,"South Sudan":90,"Spain":15,"Sri Lanka":35,"Sudan":85,"Sweden":5,"Switzerland":3,"Syria":90,"Taiwan":45,"Tajikistan":45,"Tanzania":30,"Thailand":35,"Tunisia":35,"Turkey":50,"Turkmenistan":50,"Uganda":45,"Ukraine":90,"United Arab Emirates":15,"United Kingdom":15,"United States":25,"Uruguay":10,"Uzbekistan":35,"Venezuela":75,"Vietnam":20,"Yemen":90,"Zambia":35,"Zimbabwe":55}
        return jsonify(fallback_scores)

# ── REPORT GENERATION ─────────────────────────────────────────────────────────
@app.route('/api/generate-report', methods=['POST'])
def generate_report():
    data       = request.json
    country    = data.get('country', 'Global')
    fmt        = data.get('format', 'PDF báo cáo đầy đủ')
    lang       = data.get('lang', 'Tiếng Việt')
    user_prompt = data.get('prompt', '').strip()

    now_str   = datetime.now().strftime('%d/%m/%Y %H:%M')
    timestamp = datetime.now().strftime('%Y%m%d_%H%M%S')

    # Dùng prompt của user trực tiếp, bổ sung context quốc gia
    full_prompt = (
        f"Viết báo cáo phân tích chuyên nghiệp bằng {lang} về {country}.\n"
        f"Yêu cầu cụ thể: {user_prompt}\n\n"
        f"Cấu trúc bắt buộc gồm 4 phần, mỗi phần bắt đầu bằng '## ':\n"
        f"## 1. Tổng quan\n## 2. Phân tích chi tiết\n## 3. Đánh giá rủi ro\n## 4. Khuyến nghị\n"
        f"Mỗi phần viết 4-6 câu chi tiết, thực tế, dựa trên yêu cầu của người dùng."
    )
    try:
        content_text = generate_gemini_content(full_prompt).text
    except Exception as e:
        return jsonify({"error": str(e)}), 500

    # Parse sections
    sections, current_key, current_lines = {}, None, []
    for line in content_text.splitlines():
        if line.startswith('## '):
            if current_key:
                sections[current_key] = '\n'.join(current_lines).strip()
            current_key, current_lines = line[3:].strip(), []
        else:
            current_lines.append(line)
    if current_key:
        sections[current_key] = '\n'.join(current_lines).strip()
    if not sections:
        sections = {'Nội dung': content_text}

    # ── PDF ──
    if 'PDF' in fmt:
        buf = io.BytesIO()
        doc = SimpleDocTemplate(buf, pagesize=A4,
                                leftMargin=2*cm, rightMargin=2*cm,
                                topMargin=2*cm, bottomMargin=2*cm)
        styles = getSampleStyleSheet()

        import reportlab as _rl
        _fd  = os.path.join(os.path.dirname(_rl.__file__), 'fonts')
        _wf  = 'C:/Windows/Fonts'
        def _reg(name, paths):
            for p in paths:
                if os.path.exists(p):
                    try: pdfmetrics.registerFont(TTFont(name, p)); return True
                    except: pass
            return False

        if _reg('VN', [f'{_wf}/arial.ttf']):
            _reg('VN-B', [f'{_wf}/arialbd.ttf'])
            fn, fb = 'VN', 'VN-B'
        else:
            _reg('VN', [f'{_fd}/Vera.ttf']); _reg('VN-B', [f'{_fd}/VeraBd.ttf'])
            fn, fb = 'VN', 'VN-B'

        T  = ParagraphStyle('T',  parent=styles['Title'],   fontName=fb, fontSize=18,
                            textColor=colors.HexColor('#1e3a5f'), spaceAfter=6, alignment=1)
        S  = ParagraphStyle('S',  parent=styles['Normal'],  fontName=fn, fontSize=9,
                            textColor=colors.grey, spaceAfter=12, alignment=1)
        H2 = ParagraphStyle('H2', parent=styles['Heading2'],fontName=fb, fontSize=12,
                            textColor=colors.HexColor('#1e3a5f'), spaceBefore=14, spaceAfter=6)
        B  = ParagraphStyle('B',  parent=styles['Normal'],  fontName=fn, fontSize=10,
                            leading=16, spaceAfter=8)

        story = [
            Paragraph(f"BÁO CÁO PHÂN TÍCH: {country.upper()}", T),
            Paragraph(f"Ngôn ngữ: {lang} &nbsp;|&nbsp; {now_str}", S),
            HRFlowable(width="100%", thickness=2, color=colors.HexColor('#38bdf8'), spaceAfter=12),
        ]

        for title, body in sections.items():
            story.append(Paragraph(title, H2))
            for para in re.sub(r'\*\*(.*?)\*\*', r'\1', body).split('\n'):
                if para.strip():
                    story.append(Paragraph(para.strip(), B))
            story.append(Spacer(1,6))

        story += [HRFlowable(width="100%",thickness=1,color=colors.grey,spaceBefore=12),
                  Paragraph(f"Global Pulse AI · {now_str}", S)]
        doc.build(story)
        buf.seek(0)
        return send_file(buf, mimetype='application/pdf', as_attachment=True,
                         download_name=f'report_{country}_{timestamp}.pdf')

    # ── EXCEL ──
    elif 'Excel' in fmt:
        buf = io.BytesIO()
        wb  = openpyxl.Workbook()
        ws  = wb.active
        ws.title = country[:31]
        hf = PatternFill('solid', fgColor='1e3a5f')
        ws.append(['BÁO CÁO PHÂN TÍCH', country, lang, now_str])
        for cell in ws[1]:
            cell.fill = hf; cell.font = Font(bold=True,color='FFFFFF',size=12)
            cell.alignment = Alignment(horizontal='center')
        ws.append([])
        for title, body in sections.items():
            ws.append([title])
            ws.cell(ws.max_row,1).font = Font(bold=True,size=11,color='1e3a5f')
            for line in re.sub(r'\*\*(.*?)\*\*',r'\1',body).split('\n'):
                if line.strip(): ws.append(['',line.strip()])
            ws.append([])
        ws.column_dimensions['A'].width = 25
        ws.column_dimensions['B'].width = 80
        wb.save(buf); buf.seek(0)
        return send_file(buf,
            mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
            as_attachment=True, download_name=f'report_{country}_{timestamp}.xlsx')

    # ── POWERPOINT ──
    else:
        buf = io.BytesIO()
        prs = Presentation()
        prs.slide_width  = Inches(13.33)
        prs.slide_height = Inches(7.5)
        DARK  = RGBColor(0x07,0x0a,0x0f)
        BLUE  = RGBColor(0x38,0xbd,0xf8)
        WHITE = RGBColor(0xff,0xff,0xff)
        GREY  = RGBColor(0xaa,0xbb,0xcc)

        def add_slide(ttl, bdy, cover=False):
            from pptx.util import Emu
            sl = prs.slides.add_slide(prs.slide_layouts[6])
            bg = sl.background.fill; bg.solid(); bg.fore_color.rgb = DARK
            tb = sl.shapes.add_textbox(Inches(.5),Inches(.4),Inches(12),Inches(1.2))
            p  = tb.text_frame.paragraphs[0]; p.text = ttl
            p.runs[0].font.size = Pt(28 if cover else 22)
            p.runs[0].font.bold = True; p.runs[0].font.color.rgb = BLUE
            ln = sl.shapes.add_shape(1,Inches(.5),Inches(1.7),Inches(12),Emu(40000))
            ln.fill.solid(); ln.fill.fore_color.rgb = BLUE; ln.line.fill.background()
            tb2 = sl.shapes.add_textbox(Inches(.5),Inches(2),Inches(12),Inches(5))
            tf2 = tb2.text_frame; tf2.word_wrap = True
            for i,ln_txt in enumerate(re.sub(r'\*\*(.*?)\*\*',r'\1',bdy).split('\n')):
                if not ln_txt.strip(): continue
                p2 = tf2.paragraphs[0] if i==0 else tf2.add_paragraph()
                p2.text = ln_txt.strip()
                p2.runs[0].font.size = Pt(14); p2.runs[0].font.color.rgb = WHITE
            ft = sl.shapes.add_textbox(Inches(.5),Inches(7),Inches(12),Inches(.4))
            ft.text_frame.paragraphs[0].text = f"Global Pulse AI · {now_str}"
            ft.text_frame.paragraphs[0].runs[0].font.size = Pt(9)
            ft.text_frame.paragraphs[0].runs[0].font.color.rgb = GREY

        add_slide(f"BÁO CÁO PHÂN TÍCH\n{country.upper()}",
                  f"Ngôn ngữ: {lang}\nNgày tạo: {now_str}",
                  cover=True)
        for title, body in sections.items():
            add_slide(title, body)
        prs.save(buf); buf.seek(0)
        return send_file(buf,
            mimetype='application/vnd.openxmlformats-officedocument.presentationml.presentation',
            as_attachment=True, download_name=f'report_{country}_{timestamp}.pptx')

if __name__ == '__main__':
    app.run(debug=True, port=5000)
